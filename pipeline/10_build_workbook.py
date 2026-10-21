import os, sys
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT); sys.path.insert(0, os.path.join(ROOT, "src"))  # run from anywhere; paths are relative to the repo root
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

df = pd.read_csv("data/monk_deals_simulated.csv", parse_dates=["deal_date"])
N = len(df)
cols = list(df.columns)
L = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}

F = "Arial"
hdr_font = Font(name=F, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="1F3A5F")
body = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
input_fill = PatternFill("solid", fgColor="FFF2CC")

wb = Workbook()

# ---------- README ----------
ws = wb.active
ws.title = "README"
lines = [
    ("Monk Deal Intelligence: simulated deal dataset", True),
    ("", False),
    (f"{N:,} influencer deals, {df.creator_id.nunique()} creators, {df.brand.nunique()} brands, "
     f"{df.deal_date.min().date()} to {df.deal_date.max().date()}.", False),
    ("ALL DATA IS SIMULATED. Creator and brand names are fictional (Indian regional naming for realism). Any resemblance to a real person is coincidental. Nothing here is Monk's real data.", True),
    ("Patterns are planted to mimic how real influencer data behaves, so the analysis and app can be built and tested end to end.", False),
    ("", False),
    ("Tabs", True),
    ("Deals: the main table, one row per brand deal. Filter and sort freely.", False),
    ("Creators: master list of every creator (name, handle, language, city, niche, tier, followers) with live deal count, spend, revenue and ROAS.", False),
    ("Data_Dictionary: what every column means and how it is derived.", False),
    ("Quick_Stats: live formulas on the Deals tab (spend, revenue, ROAS by tier).", False),
    ("Template: blank headers plus one example row. Paste Monk's real deals here to run the same analysis on real data.", False),
    ("", False),
    ("Required columns for real data (minimum): deal_date, creator_name, tier, brand_category, format, hook_type, "
     "duration_sec, cta_position, deal_value_inr, views, engagements, link_clicks, conversions, revenue_inr.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=F, size=12 if i == 1 else 10, bold=b)
ws.column_dimensions["A"].width = 130

# ---------- Deals ----------
wd = wb.create_sheet("Deals")
wd.append(cols)
for row in df.itertuples(index=False):
    wd.append([None if (isinstance(v, float) and pd.isna(v)) else (v.to_pydatetime() if hasattr(v, "to_pydatetime") else v)
               for v in row])
for c in wd[1]:
    c.font, c.fill = hdr_font, hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
wd.row_dimensions[1].height = 32
fmt = {"deal_date": "yyyy-mm-dd", "followers": "#,##0", "deal_value_inr": "#,##0", "views": "#,##0",
       "likes": "#,##0", "comments": "#,##0", "saves": "#,##0", "shares": "#,##0", "link_clicks": "#,##0",
       "conversions": "#,##0", "revenue_inr": "#,##0", "engagements": "#,##0", "engagement_rate": "0.0%",
       "ctr": "0.00%", "cpv_inr": "0.00", "cpe_inr": "0.00", "cost_per_conversion_inr": "#,##0", "roas": "0.00"}
widths = {"hook_line": 48, "script_text": 70, "creator_name": 18, "deal_date": 12}
for i, c in enumerate(cols, 1):
    cl = get_column_letter(i)
    wd.column_dimensions[cl].width = widths.get(c, max(11, min(len(c) + 3, 20)))
    for r in range(2, N + 2):
        cell = wd.cell(row=r, column=i)
        cell.font = body
        if c in fmt:
            cell.number_format = fmt[c]
wd.freeze_panes = "E2"
wd.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{N + 1}"

# ---------- Data dictionary ----------
dd = [
    ("deal_id", "Unique deal ID"), ("deal_date", "Date the content went live"),
    ("creator_id", "Unique creator ID"), ("creator_name", "Fictional creator name (realistic Indian regional naming, not real people)"),
    ("creator_handle", "Fictional Instagram-style handle"),
    ("creator_language", "Primary content language: Hindi, Hinglish, English, Tamil, Telugu, Kannada, Malayalam, Marathi, Gujarati, Punjabi, Bengali"),
    ("creator_city", "City the creator is based in"),
    ("platform", "instagram or youtube"), ("tier", "nano <20k, micro 20k-100k, mid 100k-500k, macro 500k-2M, mega 2M+ followers"),
    ("creator_niche", "Creator's main content niche"), ("followers", "Follower or subscriber count at time of deal"),
    ("brand", "Fictional brand name"), ("brand_category", "Category of the promoted product (can differ from creator_niche)"),
    ("format", "reel, static, story_set, youtube_integration, youtube_dedicated"),
    ("duration_sec", "Video length in seconds (blank for static and story sets)"),
    ("hook_type", "Opening style: story, question, shock, social_proof, challenge, demo, direct_pitch"),
    ("hook_line", "First line of the script (blank for non-video)"),
    ("cta_position", "Where the call to action appears: start, mid, end (blank for non-video)"),
    ("script_text", "Full short script text, used for the LLM tagging layer"),
    ("post_day", "Day of week posted"), ("post_hour", "Hour posted, 0-23 local time"),
    ("creator_deals_last_30d", "Number of deals the creator had in the 30 days before this one (saturation signal)"),
    ("deal_value_inr", "Fee paid for the deal in INR"), ("views", "Views or reach"),
    ("likes", "Likes"), ("comments", "Comments"), ("saves", "Saves"), ("shares", "Shares"),
    ("link_clicks", "Clicks on the tracked link"), ("conversions", "Purchases attributed to the deal"),
    ("revenue_inr", "Brand revenue attributed to the deal, INR"),
    ("engagements", "likes + comments + saves + shares"),
    ("engagement_rate", "engagements / views"), ("ctr", "link_clicks / views"),
    ("cpv_inr", "deal_value_inr / views"), ("cpe_inr", "deal_value_inr / engagements"),
    ("cost_per_conversion_inr", "deal_value_inr / conversions (blank if zero conversions)"),
    ("roas", "revenue_inr / deal_value_inr"),
]
assert [d[0] for d in dd] == cols, "dictionary out of sync with columns"
wk = wb.create_sheet("Data_Dictionary")
wk.append(["Column", "Meaning"])
for d in dd:
    wk.append(list(d))
for c in wk[1]:
    c.font, c.fill = hdr_font, hdr_fill
for r in wk.iter_rows(min_row=2):
    for c in r:
        c.font = body
wk.column_dimensions["A"].width = 26
wk.column_dimensions["B"].width = 95

# ---------- Quick stats (formulas) ----------
q = wb.create_sheet("Quick_Stats")
q.append(["Tier", "Deals", "Spend (INR)", "Revenue (INR)", "ROAS", "Share of spend"])
rng_ = lambda col: f"Deals!${L[col]}$2:${L[col]}${N + 1}"
tiers = ["nano", "micro", "mid", "macro", "mega"]
for i, t in enumerate(tiers, 2):
    q.append([t,
              f"=COUNTIF({rng_('tier')},A{i})",
              f"=SUMIF({rng_('tier')},A{i},{rng_('deal_value_inr')})",
              f"=SUMIF({rng_('tier')},A{i},{rng_('revenue_inr')})",
              f"=IFERROR(D{i}/C{i},0)",
              f"=IFERROR(C{i}/$C$7,0)"])
q.append(["Total", "=SUM(B2:B6)", "=SUM(C2:C6)", "=SUM(D2:D6)", "=IFERROR(D7/C7,0)", "=SUM(F2:F6)"])
for c in q[1]:
    c.font, c.fill = hdr_font, hdr_fill
for r in q.iter_rows(min_row=2, max_row=7):
    for c in r:
        c.font = bold if r[0].row == 7 else body
    r[1].number_format = "#,##0"; r[2].number_format = "#,##0"; r[3].number_format = "#,##0"
    r[4].number_format = "0.00"; r[5].number_format = "0.0%"
q["A9"] = "ROAS = revenue attributed / fee paid. Aggregate ROAS per tier, computed live from the Deals tab."
q["A9"].font = Font(name=F, size=9, italic=True)
for cl, w_ in zip("ABCDEF", [12, 10, 16, 16, 10, 16]):
    q.column_dimensions[cl].width = w_

# ---------- Creators master ----------
cm = (df.sort_values("deal_date").groupby("creator_id").agg(
    creator_name=("creator_name", "first"), creator_handle=("creator_handle", "first"),
    creator_language=("creator_language", "first"), creator_city=("creator_city", "first"),
    creator_niche=("creator_niche", "first"), platform=("platform", "first"), tier=("tier", "first"),
    followers=("followers", "last")).reset_index().sort_values("creator_id"))
ccols = ["creator_id", "creator_name", "creator_handle", "creator_language", "creator_city", "creator_niche",
         "platform", "tier", "followers", "total_deals", "total_spend_inr", "total_revenue_inr", "roas"]
wc = wb.create_sheet("Creators")
wc.append(ccols)
for i, r in enumerate(cm.itertuples(index=False), 2):
    wc.append([r.creator_id, r.creator_name, r.creator_handle, r.creator_language, r.creator_city, r.creator_niche,
               r.platform, r.tier, int(r.followers),
               f"=COUNTIF({rng_('creator_id')},A{i})",
               f"=SUMIF({rng_('creator_id')},A{i},{rng_('deal_value_inr')})",
               f"=SUMIF({rng_('creator_id')},A{i},{rng_('revenue_inr')})",
               f"=IFERROR(L{i}/K{i},0)"])
for c in wc[1]:
    c.font, c.fill = hdr_font, hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for r in wc.iter_rows(min_row=2):
    for c in r:
        c.font = body
    r[8].number_format = "#,##0"; r[9].number_format = "#,##0"; r[10].number_format = "#,##0"
    r[11].number_format = "#,##0"; r[12].number_format = "0.00"
for i, w_ in enumerate([11, 24, 28, 14, 16, 12, 11, 8, 12, 12, 16, 16, 8], 1):
    wc.column_dimensions[get_column_letter(i)].width = w_
wc.freeze_panes = "C2"
wc.auto_filter.ref = f"A1:{get_column_letter(len(ccols))}{len(cm) + 1}"

# ---------- Template ----------
t = wb.create_sheet("Template")
t.append(cols)
ex = {"deal_id": "D9001", "deal_date": "2026-10-01", "creator_id": "C999", "creator_name": "Example Creator",
      "creator_handle": "@example.creator", "creator_language": "Hinglish", "creator_city": "Mumbai",
      "platform": "instagram", "tier": "mid", "creator_niche": "beauty", "followers": 250000, "brand": "Example Brand",
      "brand_category": "beauty", "format": "reel", "duration_sec": 30, "hook_type": "story",
      "hook_line": "Last month my skin broke out before a shoot.", "cta_position": "mid",
      "script_text": "Paste full script here", "post_day": "Thursday", "post_hour": 20,
      "creator_deals_last_30d": 1, "deal_value_inr": 35000, "views": 62000, "likes": 1800, "comments": 90,
      "saves": 300, "shares": 80, "link_clicks": 1100, "conversions": 55, "revenue_inr": 40000,
      "engagements": 2270, "engagement_rate": 0.0366, "ctr": 0.0177, "cpv_inr": 0.56, "cpe_inr": 15.4,
      "cost_per_conversion_inr": 636, "roas": 1.14}
t.append([ex[c] for c in cols])
for c in t[1]:
    c.font, c.fill = hdr_font, hdr_fill
    c.alignment = Alignment(horizontal="center", wrap_text=True)
for c in t[2]:
    c.font, c.fill = body, input_fill
t["A4"] = "Row 2 (yellow) is an example only. Delete it, then paste real deals from row 2 down. Derived columns (engagements to roas) can be left blank; the app will compute them."
t["A4"].font = Font(name=F, size=9, italic=True)
for i, c in enumerate(cols, 1):
    t.column_dimensions[get_column_letter(i)].width = widths.get(c, max(11, min(len(c) + 3, 20)))
t.freeze_panes = "A2"

wb.save("docs/Monk_Deal_Intelligence_Simulated_Data.xlsx")
print("saved", N, "rows")
